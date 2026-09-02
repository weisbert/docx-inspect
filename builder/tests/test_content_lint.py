#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Tests for content_lint.py.

A single counter-example fixture project deliberately trips EVERY finding type,
so the linter's coverage is asserted rather than assumed (the golden render only
ever exercised missing_image). Beyond coverage this file pins:

  * the Check-tab payload: grouped errors / warnings / notes, and the per-item
    keys level / code / loc / nodeId / blockId / message,
  * the flat-list compatibility the export path relies on,
  * the sim_span rule matching what tables.py actually merges -- cross-checked
    against the REAL renderer, not against a restatement of it,
  * every rule paired: it fires on a fixture built to trip it, and stays silent
    on the same fixture with the mistake corrected,
  * a blank image file being a NOTE, so an inherited report full of empty frames
    does not read as broken,
  * the CLI exit codes (0 clean / 1 any error / 2 no project.json).

Run:

    python builder/tests/test_content_lint.py
"""
import contextlib
import io
import json
import os
import re
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))  # builder/tests
BUILDER = os.path.dirname(HERE)                  # builder/
if BUILDER not in sys.path:
    sys.path.insert(0, BUILDER)
import buildpath  # noqa: E402,F401  (registers core/docx_io/store/sync/web)

import content_lint as cl  # noqa: E402

# A project crafted so each rule fires at least once. Comments name the target.
FIXTURE = {
    "schema_version": 1,
    "outline": [
        {"id": "n1", "title": "Empty", "blocks": [], "children": []},   # empty_section
        {"id": "n2", "title": "Media", "blocks": [
            {"type": "image", "id": "i1", "file": "C:/abs/x.png", "caption": ""},   # image_path(abs)+no_caption
            {"type": "image", "id": "i2", "file": "../up.png", "caption": "ok"},     # image_path(..)
            {"type": "image", "id": "i3", "file": "notimages/x.png", "caption": "ok"},  # image_path(not images/)
            {"type": "image", "id": "i4", "file": "", "caption": "cleared"},         # image_placeholder (note)
            {"type": "imagegrid", "id": "g1",
             "items": [{"file": "/abs.png"}, {"file": ""}], "caption": ""},          # image_path+image_placeholder+no_caption
            {"type": "table", "id": "t1"},                                          # table_no_rows
            {"type": "table", "id": "t2", "rows": [["a", "b"], ["c", "d"]],
             "merges": [{"r": 0, "c": 0, "rs": 5, "cs": 1}],
             "row_fills": {"9": "FFFF00"}},                                          # free_table_bounds x2
            {"type": "datatable", "id": "d0"},                                       # datatable_no_data
        ], "children": [
            {"id": "n2a", "title": "Nested empty", "blocks": [], "children": []},     # empty_section at loc 2.1
        ]},
        {"id": "n3", "title": "Compliance", "blocks": [
            {"type": "datatable", "id": "d1", "caption": "C", "data": {
                "spec_name": "Spec",
                # two groups of two axes each: BOTH merge fine, so a sim_span row
                # here must NOT be reported (see test_sim_span_matches_renderer)
                "sims": [{"key": "s1", "title": "S1", "axes": ["MIN", "TYP"]},
                         {"key": "s2", "title": "S2", "axes": ["MIN", "TYP"]}],
                "rows": [
                    {"item": "noKeys"},                                              # row_missing_key
                    {"cat": "C", "item": "r1", "kind": "result", "unit": "",         # unit_empty(note)
                     "limit": None, "spec": 5, "spec_mtm": [None, 5, None],          # limit_no_flag
                     "sim_span": True, "sims": {"bogus": {"mtm": [1, 2, 3]}}},        # unknown_sim_key + sim_span
                    {"cat": "C", "item": "r2", "kind": "result", "unit": "uA",
                     "limit": "le", "spec_mtm": [None, None, None],
                     "sim_mtm": [None, None, None]},                                  # empty_sim_result
                ]}},
            {"type": "datatable", "id": "d2", "caption": "C2", "data": {
                "sims": [{"key": "s1", "axes": ["MIN", "TYP", "MAX", "NTWC"]}],
                "rows": [{"cat": "X", "item": "y", "kind": "result", "unit": "V",
                          "limit": "le", "sim_mtm": [1, 2, 3]}]}},                    # no_setting_rows
            {"type": "datatable", "id": "d3", "caption": "C3", "data": {
                # a one-axis group cannot merge -- the renderer says so too
                "sims": [{"key": "s1", "axes": ["TYP"]}],
                "rows": [
                    {"cat": "Cond", "item": "temp", "kind": "common_setting",
                     "unit": "C", "sims": {"s1": {"mtm": [25]}}},
                    {"cat": "X", "item": "z", "kind": "result", "unit": "V",
                     "limit": "le", "sim_span": True,
                     "sims": {"s1": {"mtm": [1]}}},
                ]}},                                                                  # sim_span_unmergeable
        ], "children": []},
    ],
}

EXPECTED = {
    # every code content_lint can emit from this fixture, with its level
    "empty_section": "warn",
    "image_path": "error",
    "image_placeholder": "info",
    "no_caption": "warn",
    "table_no_rows": "error",
    "free_table_bounds": "error",
    "datatable_no_data": "error",
    "row_missing_key": "error",
    "sim_span_unmergeable": "error",
    "limit_no_flag": "warn",
    "unknown_sim_key": "warn",
    "empty_sim_result": "warn",
    "no_setting_rows": "warn",
    "unit_empty": "info",
}

fails = 0
skipped = []


def check(cond, name, detail=""):
    global fails
    print(("  PASS " if cond else "  FAIL ") + name + ("" if cond else "  -> " + detail))
    if not cond:
        fails += 1


def skip(name, why):
    """Report a check that could NOT run. Never counted as a pass -- a silent
    skip is how a cross-check quietly stops guarding anything."""
    skipped.append(name)
    print("  SKIP " + name + "  -> " + why)


def _is_cjk(ch):
    cp = ord(ch)
    return (0x3400 <= cp <= 0x4DBF or 0x4E00 <= cp <= 0x9FFF
            or 0xF900 <= cp <= 0xFAFF or 0x20000 <= cp <= 0x2FA1F)


def test_coverage(report):
    got = {}
    for f in report.flat:
        got.setdefault(f["code"], []).append(f)

    # every expected code fires...
    for t, lv in EXPECTED.items():
        check(t in got, "fires: %s" % t, "not found; got=%r" % sorted(got))
        if t in got:
            check(all(f["level"] == lv for f in got[t]),
                  "level(%s)==%s" % (t, lv),
                  "levels=%r" % [f["level"] for f in got[t]])

    # ...and nothing unexpected fires (guards against a rule going rogue)
    unexpected = set(got) - set(EXPECTED)
    check(not unexpected, "no unexpected finding codes", "extra=%r" % sorted(unexpected))

    # count-sensitive spots
    check(len(got.get("image_path", [])) == 4, "4 image_path findings",
          "got %d" % len(got.get("image_path", [])))
    check(len(got.get("image_placeholder", [])) == 2, "2 image_placeholder notes",
          "got %d" % len(got.get("image_placeholder", [])))
    check(len(got.get("free_table_bounds", [])) == 2, "2 free_table_bounds findings",
          "got %d" % len(got.get("free_table_bounds", [])))
    check(len(got.get("no_setting_rows", [])) == 2, "2 no_setting_rows findings",
          "got %d" % len(got.get("no_setting_rows", [])))
    check(len(got.get("empty_section", [])) == 2, "2 empty_section findings",
          "got %d" % len(got.get("empty_section", [])))
    return got


def test_item_shape(report, got):
    keys = {"type", "level", "detail", "location", "code", "message",
            "loc", "nodeId", "blockId"}
    check(all(keys <= set(f) for f in report.flat), "every item has the full key set",
          "%r" % [sorted(set(f)) for f in report.flat[:1]])
    check(all(f["type"] == f["code"] and f["detail"] == f["message"]
              for f in report.flat),
          "legacy type/detail alias code/message")

    # section number path
    check(got["empty_section"][0]["loc"] == "1", "loc of the first section is '1'",
          "%r" % got["empty_section"][0]["loc"])
    check(got["empty_section"][1]["loc"] == "2.1", "nested section loc is '2.1'",
          "%r" % got["empty_section"][1]["loc"])
    check(all(f["loc"] == "3" for f in got["no_setting_rows"]),
          "compliance findings carry loc '3'",
          "%r" % [f["loc"] for f in got["no_setting_rows"]])

    # node / block ids for the jump-to-block behaviour
    check(got["empty_section"][0]["nodeId"] == "n1" and
          got["empty_section"][0]["blockId"] is None,
          "section finding carries nodeId and a null blockId")
    check(all(f["nodeId"] == "n3" for f in got["no_setting_rows"]),
          "block finding carries its section's nodeId")
    check(sorted(f["blockId"] for f in got["no_setting_rows"]) == ["d1", "d2"],
          "block finding carries its blockId",
          "%r" % sorted(f["blockId"] for f in got["no_setting_rows"]))
    check(sorted(f["blockId"] for f in got["image_path"]) == ["g1", "i1", "i2", "i3"],
          "image findings carry their blockId",
          "%r" % sorted(f["blockId"] for f in got["image_path"]))

    # messages are plain English, no Chinese anywhere in the frozen table
    bad = [m for m in cl.MESSAGES.values() if any(_is_cjk(c) for c in m)]
    check(not bad, "no CJK in the frozen message table", "%r" % bad)
    check(all(f["message"] and not any(_is_cjk(c) for c in f["message"])
              for f in report.flat), "no CJK in any emitted message")
    check(all(f["message"] != f.get("code") for f in report.flat),
          "every emitted message resolved to a real sentence (no bare id)",
          "%r" % [f["message"] for f in report.flat if f["message"] == f["code"]])
    check("never marked over spec" in got["limit_no_flag"][0]["message"],
          "blank-limit wording matches the glossary",
          got["limit_no_flag"][0]["message"])
    check("setting rows" in got["no_setting_rows"][0]["message"],
          "missing-setting-rows wording names setting rows",
          got["no_setting_rows"][0]["message"])


def test_grouping(report):
    check(sorted(report.keys()) == ["errors", "notes", "warnings"],
          "report has exactly errors / warnings / notes", "%r" % sorted(report.keys()))
    n = len(report["errors"]) + len(report["warnings"]) + len(report["notes"])
    check(n == len(report.flat), "groups partition the flat list",
          "%d vs %d" % (n, len(report.flat)))
    check(all(f["level"] == "error" for f in report["errors"])
          and all(f["level"] == "warn" for f in report["warnings"])
          and all(f["level"] == "info" for f in report["notes"]),
          "each group holds only its own level")
    check(report.has_errors, "has_errors is true for the fixture")
    check(report.counts == cl.summarize(report.flat), "counts matches summarize()")

    # the Check tab reads this straight off the wire
    payload = json.loads(json.dumps(report))
    check(isinstance(payload, dict) and sorted(payload) == ["errors", "notes", "warnings"],
          "json.dumps(report) is the grouped object", "%r" % type(payload))
    check(payload["errors"][0]["code"] in EXPECTED, "json items keep their code")

    # the export path merges the findings as a flat list: list(report) + list(w)
    merged = list(report) + [{"type": "missing_logo", "level": "warn"}]
    check(len(merged) == len(report.flat) + 1, "list(report) yields the flat findings",
          "%d" % len(merged))
    check(merged[0] is report.flat[0], "flat order is document order")


def test_sim_span_matches_renderer():
    """tables.py merges each sim group's OWN axis columns (cc[0]..cc[-1]), so a
    group needs >= 2 axis columns and several groups are perfectly fine. The old
    test asserted the pre-fix rule (>= 3 axes, single group only); the renderer
    moved on and the linter must mirror the renderer, not the old doc."""
    def rows(span):
        return [{"cat": "c", "item": "i", "kind": "result", "unit": "V",
                 "limit": "le", "sim_span": span, "sims": {"s1": {"mtm": [1, 2]}}},
                {"cat": "c", "item": "s", "kind": "common_setting", "unit": "V"}]

    def lint(sims, span=True):
        p = {"outline": [{"id": "n", "title": "S", "blocks": [
            {"type": "datatable", "id": "b", "data": {"sims": sims, "rows": rows(span)}}],
            "children": []}]}
        return [f["code"] for f in cl.lint_project(p, {}).flat]

    two_axes = [{"key": "s1", "axes": ["MIN", "TYP"]}]
    check("sim_span_unmergeable" not in lint(two_axes),
          "sim_span over a 2-axis group is fine", "%r" % lint(two_axes))
    two_groups = [{"key": "s1", "axes": ["MIN", "TYP", "MAX"]},
                  {"key": "s2", "axes": ["MIN", "TYP", "MAX"]}]
    check("sim_span_unmergeable" not in lint(two_groups),
          "sim_span across two full groups is fine", "%r" % lint(two_groups))
    thin = [{"key": "s1", "axes": ["MIN", "TYP", "MAX"]}, {"key": "s2", "axes": ["TYP"]}]
    check("sim_span_unmergeable" in lint(thin),
          "a single thin group is caught even when the others are wide",
          "%r" % lint(thin))
    check("sim_span_unmergeable" not in lint(thin, span=False),
          "no sim_span row -> no complaint about thin groups", "%r" % lint(thin, False))
    check(cl.LEVELS["sim_span_unmergeable"] == "error",
          "sim_span_unmergeable is an error (the renderer drops the merge)")


def test_missing_image_needs_a_folder():
    """The blank-file note and the missing-file warning are different states, and
    the missing-file check only runs when a report folder is supplied (the export
    path does not supply one -- the renderer reports missing pictures itself)."""
    tmp = tempfile.mkdtemp(prefix="lint_files_")
    try:
        os.makedirs(os.path.join(tmp, "images"))
        with open(os.path.join(tmp, "images", "there.png"), "wb") as fh:
            fh.write(b"\x89PNG\r\n")
        proj = {"outline": [{"id": "n", "title": "S", "blocks": [
            {"type": "image", "id": "a", "file": "images/there.png", "caption": "c"},
            {"type": "image", "id": "b", "file": "images/gone.png", "caption": "c"},
        ], "children": []}]}
        no_dir = [f["code"] for f in cl.lint_project(proj, {}).flat]
        check(no_dir == [], "without a folder there is nothing to report", "%r" % no_dir)
        with_dir = cl.lint_project(proj, {}, tmp).flat
        codes = [f["code"] for f in with_dir]
        check(codes == ["missing_image"], "with a folder the missing file is reported",
              "%r" % codes)
        check(with_dir[0]["blockId"] == "b", "the missing-file finding points at block b")
        check(with_dir[0]["level"] == "warn", "a missing picture is a warning")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_helpers_and_edge_cases():
    check(cl.classify("block_error") == "error", "classify block_error -> error")
    check(cl.classify("totally_new_code") == "warn", "classify unknown -> warn")
    ws = [{"type": "missing_image"}, {"type": "block_error"}, {"type": "x", "level": "info"}]
    cl.stamp_levels(ws)
    check(ws[0]["level"] == "warn" and ws[1]["level"] == "error" and ws[2]["level"] == "info",
          "stamp_levels stamps + preserves existing")
    check(cl.message("no_caption.image") == "Image has no caption",
          "message() renders a frozen string")
    check(cl.message("not.a.real.id") == "not.a.real.id",
          "message() degrades to the id instead of raising")

    # a clean project yields nothing, in both shapes
    clean = {"schema_version": 1, "outline": [
        {"id": "n", "title": "S", "blocks": [{"type": "para", "runs": [{"t": "hi"}]}],
         "children": []}]}
    rep = cl.lint_project(clean, {})
    check(rep.flat == [] and list(rep) == [], "clean project -> no findings", "%r" % rep.flat)
    check(rep == {"errors": [], "warnings": [], "notes": []},
          "clean project -> three empty groups", "%r" % dict(rep))
    check(not rep.has_errors, "clean project has no errors")

    # regression: a malformed SCALAR sim_mtm / spec_mtm must NOT crash the lint
    # (a lint gate that blows up on the malformed input it exists to catch is worse
    # than useless). It should tolerate it and still emit findings for the section.
    bad = {"schema_version": 1, "outline": [{"title": "S", "blocks": [
        {"type": "datatable", "data": {"sims": [{"key": "s1"}], "rows": [
            {"cat": "c", "item": "i", "kind": "result", "unit": "V",
             "sim_mtm": 5, "spec_mtm": 9, "limit": "le"}]}}], "children": []}]}
    try:
        cl.lint_project(bad, {})
        check(True, "scalar sim_mtm/spec_mtm does not crash the lint")
    except Exception as e:
        check(False, "scalar sim_mtm/spec_mtm does not crash the lint", repr(e))

    # a node with no id at all still lints (nodeId is simply None)
    check(cl.lint_project(bad, {}).flat[0]["nodeId"] is None,
          "a node without an id yields nodeId None")


# ---------------------------------------------------------------------------
# 1. The two cases that used to fail, and WHY they were wrong.
# ---------------------------------------------------------------------------
# The two failures in this file before this round were `fires: sim_span_axes`
# and `fires: sim_span_multi`. Neither was a linter bug that got fixed: commit
# 0e1316e changed what tables.py DOES. It used to merge a sim_span row across
# one single block of columns starting at cc[2], so a table needed one sim group
# of at least three axes and mis-merged anything else; it now merges each sim
# group's own axis columns independently (cc[0]..cc[-1], per group). After that
# change a two-axis group merges cleanly and several groups are perfectly fine,
# so the two codes for "too few axes" and "more than one group" could never be
# emitted again -- the assertions, not the module, encoded the stale rule.
# The replacement is not another restatement of the rule: the linter is checked
# AGAINST THE RENDERER ITSELF below, so the two can never drift apart again.

_COMP_CFG = {
    "col_w_cm": {"cat": 2.0, "item": 3.0, "spec": 1.5, "axis": 1.4,
                 "spacer": 0.2, "unit": 1.2},
    "font_pt": 7,
    "row_h_pt": {"header": 12, "data": 10},
    "axis_labels": ["MIN", "TYP", "MAX", "NTWC"],
    "fills": {"header": "FFF2CC", "setting": "DDEBF7", "result": "FFFFFF",
              "separator": "BFBFBF"},
    "setting_kinds": ["common_setting", "module_setting", "tb"],
    "flag_color": "FF0000",
    "borders": {"val": "single", "sz": 4, "color": "000000"},
}
_FULL_CFG = {"compliance": _COMP_CFG}


def _span_table_data(sims, span=True):
    """A compliance table whose result row spans its simulation columns."""
    data = {"show_spec": False,
            "rows": [
                {"cat": "Conditions", "item": "Temperature", "kind": "common_setting",
                 "unit": "C", "sim_mtm": [25, None, None]},
                {"cat": "CLKDIV_5G", "item": "Corner", "kind": "result", "unit": "",
                 "limit": "le", "sim_span": span, "sim_mtm": ["TT", None, None]},
            ]}
    if sims is not None:
        data["sims"] = sims
    return data


def _lint_codes(project, cfg=None, project_dir=None):
    return [f["code"] for f in cl.lint_project(project, cfg or {}, project_dir).flat]


def _one_block_project(block, node_id="n", title="Section"):
    return {"schema_version": 1,
            "outline": [{"id": node_id, "title": title, "blocks": [block],
                         "children": []}]}


def test_dead_codes_are_gone():
    """The two stale codes must not merely stop firing -- they must be gone from
    the module, or a future edit will quietly resurrect the pre-0e1316e rule."""
    src = open(cl.__file__, encoding="utf-8").read()
    for dead in ("sim_span_axes", "sim_span_multi"):
        check(dead not in cl.LEVELS, "%s is gone from LEVELS" % dead)
        check(dead not in src, "%s appears nowhere in content_lint.py" % dead)
    # the surviving code is the renderer's own, so a manifest never carries two
    # names for one problem
    check("sim_span_unmergeable" in cl.LEVELS
          and cl.LEVELS["sim_span_unmergeable"] == "error",
          "sim_span_unmergeable is the one surviving sim_span code, an error")
    tsrc = open(os.path.join(BUILDER, "core", "tables.py"), encoding="utf-8").read()
    check('"sim_span_unmergeable"' in tsrc,
          "tables.py emits that same code (one name, two emitters)")


def test_sim_span_mirrors_the_real_renderer():
    """Run the SAME tables through content_lint and through tables.py, and demand
    they agree on sim_span_unmergeable. This is what the two removed assertions
    should have been: a comparison with the renderer, not with the doc."""
    try:
        from docx import Document
        import tables
    except Exception as e:                       # python-docx absent
        skip("renderer cross-check", "python-docx unavailable: %s" % e)
        return

    cases = [
        ("one group of 3 axes", [{"key": "a", "axes": ["MIN", "TYP", "MAX"]}],
         True, False),
        ("one group of 2 axes", [{"key": "a", "axes": ["MIN", "TYP"]}],
         True, False),
        ("two full groups", [{"key": "a", "axes": ["MIN", "TYP", "MAX"]},
                             {"key": "b", "axes": ["MIN", "TYP", "MAX"]}],
         True, False),
        ("one thin group beside a wide one",
         [{"key": "a", "axes": ["MIN", "TYP", "MAX"]}, {"key": "b", "axes": ["TYP"]}],
         True, True),
        ("the only group is thin", [{"key": "a", "axes": ["TYP"]}],
         True, True),
        ("thin group, no spanning row",
         [{"key": "a", "axes": ["MIN", "TYP", "MAX"]}, {"key": "b", "axes": ["TYP"]}],
         False, False),
        ("no sims declared (implicit group takes the template axes)", None,
         True, False),
    ]
    for label, sims, span, expected in cases:
        data = _span_table_data(sims, span)
        rendered = tables.render_datatable(Document(), data, _COMP_CFG)
        r_fires = any(w.get("type") == "sim_span_unmergeable"
                      for w in rendered["warnings"])
        codes = _lint_codes(_one_block_project(
            {"type": "datatable", "id": "b", "data": data}), _FULL_CFG)
        l_fires = "sim_span_unmergeable" in codes
        check(r_fires == expected,
              "renderer: %s -> %s" % (label, "warns" if expected else "silent"),
              "renderer warnings=%r" % rendered["warnings"])
        check(l_fires == r_fires,
              "linter agrees with the renderer: %s" % label,
              "renderer=%s linter=%s codes=%r" % (r_fires, l_fires, codes))


# ---------------------------------------------------------------------------
# 2. Machine keys on EVERY item, on a real report shape.
# ---------------------------------------------------------------------------

_LOC_RE = re.compile(r"^\d+(\.\d+)*$")


def test_every_item_carries_loc_node_message(report):
    """The Check tab renders `loc` as the section number and jumps by nodeId /
    blockId, so a single item missing one of them is a dead row in the list."""
    flat = report.flat
    check(bool(flat), "the fixture produced findings to inspect")
    bad_loc = [f for f in flat if not _LOC_RE.match(f.get("loc") or "")]
    check(not bad_loc, "every item's loc is a section number path",
          "%r" % [(f["code"], f.get("loc")) for f in bad_loc])
    bad_msg = [f for f in flat
               if not isinstance(f.get("message"), str) or not f["message"].strip()]
    check(not bad_msg, "every item has a non-empty message",
          "%r" % [f["code"] for f in bad_msg])
    check(all("nodeId" in f for f in flat), "every item has a nodeId key")
    bad_node = [f for f in flat if f["nodeId"] not in ("n1", "n2", "n2a", "n3")]
    check(not bad_node, "every nodeId is one of the fixture's real nodes",
          "%r" % [(f["code"], f["nodeId"]) for f in bad_node])
    known_blocks = {None, "i1", "i2", "i3", "i4", "g1", "t1", "t2", "d0",
                    "d1", "d2", "d3"}
    bad_block = [f for f in flat if f["blockId"] not in known_blocks]
    check(not bad_block, "every blockId is None or one of the fixture's blocks",
          "%r" % [(f["code"], f["blockId"]) for f in bad_block])
    check(all(f.get("level") in cl.LEVEL_GROUP for f in flat),
          "every item's level is one the three groups can hold",
          "%r" % sorted({f.get("level") for f in flat}))

    # the shape the Check tab actually receives, and the shape the export path
    # actually merges -- both off the SAME object.
    payload = json.loads(json.dumps(report))
    check(sorted(payload) == ["errors", "notes", "warnings"]
          and sum(len(payload[k]) for k in payload) == len(flat),
          "the wire payload is the three groups and loses nothing",
          "%r" % {k: len(v) for k, v in payload.items()})
    merged = list(report) + [{"type": "missing_logo", "level": "warn"}]
    check(all(isinstance(w, dict) for w in merged),
          "the export merge sees only dicts (a plain dict would inject 3 strings)")
    check(sum(1 for w in merged if w.get("level") == "error")
          == len(report["errors"]),
          "the export stats loop counts the same errors the Check tab groups")


# ---------------------------------------------------------------------------
# 3. Every rule paired: it fires on the mistake, and is silent once corrected.
# ---------------------------------------------------------------------------

_GOOD_ROW = {"cat": "CLKDIV_5G", "item": "Supply current", "kind": "result",
             "unit": "uA", "limit": "le", "spec_mtm": [None, None, 500],
             "sim_mtm": [1, 2, 3]}
_GOOD_SETTING = {"cat": "Conditions", "item": "Temperature",
                 "kind": "common_setting", "unit": "C", "sim_mtm": [25, None, None]}


def _dt(rows, sims=None):
    data = {"rows": rows}
    if sims is not None:
        data["sims"] = sims
    return {"type": "datatable", "id": "b", "data": data}


def _row(**kw):
    r = dict(_GOOD_ROW)
    r.update(kw)
    return r


def _clean_table_rows(extra=None):
    return [dict(_GOOD_SETTING), dict(extra or _GOOD_ROW)]


# (code, project that trips it, project with the SAME content corrected)
_PAIRS = [
    ("empty_section",
     {"outline": [{"id": "n", "title": "Empty", "blocks": [], "children": []}]},
     {"outline": [{"id": "n", "title": "Empty",
                   "blocks": [{"type": "para", "runs": [{"t": "text"}]}],
                   "children": []}]}),
    ("image_path",
     _one_block_project({"type": "image", "id": "b", "file": "pics/a.png",
                         "caption": "Block diagram"}),
     _one_block_project({"type": "image", "id": "b", "file": "images/a.png",
                         "caption": "Block diagram"})),
    ("image_placeholder",
     _one_block_project({"type": "image", "id": "b", "file": "",
                         "caption": "Block diagram"}),
     _one_block_project({"type": "image", "id": "b", "file": "images/a.png",
                         "caption": "Block diagram"})),
    ("no_caption",
     _one_block_project({"type": "image", "id": "b", "file": "images/a.png",
                         "caption": "  "}),
     _one_block_project({"type": "image", "id": "b", "file": "images/a.png",
                         "caption": "Block diagram"})),
    ("table_no_rows",
     _one_block_project({"type": "table", "id": "b"}),
     _one_block_project({"type": "table", "id": "b", "rows": [["Pin", "Type"]]})),
    ("free_table_bounds",
     _one_block_project({"type": "table", "id": "b", "rows": [["a", "b"], ["c", "d"]],
                         "merges": [{"r": 0, "c": 0, "rs": 9, "cs": 1}]}),
     _one_block_project({"type": "table", "id": "b", "rows": [["a", "b"], ["c", "d"]],
                         "merges": [{"r": 0, "c": 0, "rs": 2, "cs": 1}]})),
    ("datatable_no_data",
     _one_block_project({"type": "datatable", "id": "b"}),
     _one_block_project(_dt(_clean_table_rows()))),
    ("row_missing_key",
     _one_block_project(_dt([dict(_GOOD_SETTING),
                             {"cat": "CLKDIV_5G", "item": "I_dd", "kind": "result"}])),
     _one_block_project(_dt(_clean_table_rows()))),
    ("no_setting_rows",
     _one_block_project(_dt([dict(_GOOD_ROW)])),
     _one_block_project(_dt(_clean_table_rows()))),
    ("limit_no_flag",
     _one_block_project(_dt(_clean_table_rows(_row(limit=None)))),
     _one_block_project(_dt(_clean_table_rows(_row(limit="le"))))),
    ("empty_sim_result",
     _one_block_project(_dt(_clean_table_rows(
         _row(sim_mtm=[None, None, None], sim_ntwc=None)))),
     _one_block_project(_dt(_clean_table_rows()))),
    ("unit_empty",
     _one_block_project(_dt(_clean_table_rows(_row(unit="")))),
     _one_block_project(_dt(_clean_table_rows()))),
    ("unknown_sim_key",
     _one_block_project(_dt(_clean_table_rows(_row(sims={"nope": {"mtm": [1, 2, 3]}})),
                            sims=[{"key": "fdr", "axes": ["MIN", "TYP", "MAX"]}])),
     _one_block_project(_dt(_clean_table_rows(_row(sims={"fdr": {"mtm": [1, 2, 3]}})),
                            sims=[{"key": "fdr", "axes": ["MIN", "TYP", "MAX"]}]))),
    ("sim_span_unmergeable",
     _one_block_project(_dt(_clean_table_rows(_row(sim_span=True)),
                            sims=[{"key": "fdr", "axes": ["TYP"]}])),
     _one_block_project(_dt(_clean_table_rows(_row(sim_span=True)),
                            sims=[{"key": "fdr", "axes": ["MIN", "TYP"]}]))),
]


def test_each_rule_fires_and_stays_silent():
    seen = set()
    for code, dirty, clean in _PAIRS:
        seen.add(code)
        d = _lint_codes(dirty, _FULL_CFG)
        c = _lint_codes(clean, _FULL_CFG)
        check(code in d, "%s fires on the mistake" % code, "codes=%r" % d)
        check(code not in c, "%s is silent once corrected" % code, "codes=%r" % c)
    # no rule may exist without a pair, or a new rule ships untested
    check(seen == set(EXPECTED),
          "every rule the fixture can produce has a fires/silent pair",
          "unpaired=%r" % sorted(set(EXPECTED) - seen))

    # the fully corrected report is clean in ALL respects, not just per-rule
    spotless = {"schema_version": 1, "outline": [
        {"id": "n1", "title": "MODULE_A overview", "blocks": [
            {"type": "para", "runs": [{"t": "The block divides a 5 GHz input."}]},
            {"type": "image", "id": "i", "file": "images/topology.png",
             "caption": "Topology"},
            _dt(_clean_table_rows()),
        ], "children": []}]}
    codes = _lint_codes(spotless, _FULL_CFG)
    check(codes == [], "a corrected report produces no findings at all",
          "codes=%r" % codes)

    # the template config calibrates the rules: a setting kind the template
    # declares must satisfy the setting-rows rule (policy lives in the template,
    # mechanism lives here).
    custom = _one_block_project(_dt([
        {"cat": "Conditions", "item": "Load", "kind": "bench_setting", "unit": "pF"},
        dict(_GOOD_ROW)]))
    cfg = {"compliance": dict(_COMP_CFG, setting_kinds=["bench_setting"])}
    check("no_setting_rows" not in _lint_codes(custom, cfg),
          "a setting kind declared by the template counts as a setting row",
          "%r" % _lint_codes(custom, cfg))
    check("no_setting_rows" in _lint_codes(custom, _FULL_CFG),
          "the same table fails against a template that does not declare it")


# ---------------------------------------------------------------------------
# 4. A blank image file is a NOTE. An inherited report is not a broken one.
# ---------------------------------------------------------------------------


def test_inherited_report_full_of_empty_frames_is_not_broken():
    """A report inherited from the previous stage keeps its figure frames and
    captions with the file cleared. Fifty empty frames must read as fifty notes:
    zero errors, zero warnings, and nothing the Check tab paints as a problem."""
    blocks = [{"type": "image", "id": "f%d" % i, "file": "",
               "caption": "Result %d" % i} for i in range(50)]
    blocks.append({"type": "para", "runs": [{"t": "Results to be filled in."}]})
    proj = {"schema_version": 1, "outline": [
        {"id": "n", "title": "Simulation Results", "blocks": blocks, "children": []}]}
    rep = cl.lint_project(proj, _FULL_CFG)
    check(rep.counts == {"error": 0, "warn": 0, "info": 50},
          "50 cleared figures -> 50 notes, no errors, no warnings", "%r" % rep.counts)
    check(rep["errors"] == [] and rep["warnings"] == [],
          "the errors and warnings groups are empty")
    check(len(rep["notes"]) == 50 and
          all(f["code"] == "image_placeholder" for f in rep["notes"]),
          "all 50 land in Notes as image_placeholder")
    check(not rep.has_errors, "an inherited report does not block an export")
    check(cl.LEVELS["image_placeholder"] == "info",
          "image_placeholder is declared info, not warn")

    # a cleared frame is NOT a missing file, even when the folder is inspected:
    # the two states must not be conflated or every inherited report reads as 50
    # warnings.
    tmp = tempfile.mkdtemp(prefix="lint_blank_")
    try:
        os.makedirs(os.path.join(tmp, "images"))
        rep2 = cl.lint_project(proj, _FULL_CFG, tmp)
        codes = {f["code"] for f in rep2.flat}
        check(codes == {"image_placeholder"},
              "with a folder, a cleared frame is still only a note",
              "%r" % sorted(codes))
        check(rep2.counts["warn"] == 0, "no missing_image for a blank file")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    # the caption survives the clearing, so the frame must not also be scolded
    # for having no caption
    check(all(f["code"] != "no_caption" for f in rep.flat),
          "a cleared frame that kept its caption is not reported as uncaptioned")


# ---------------------------------------------------------------------------
# 4a. A cross-reference whose target was deleted.
#
# Deleting a figure leaves every sentence that pointed at it holding an id that
# resolves to nothing. The renderer writes a red "[ref: <id>]" marker into that
# sentence, so the exported report says something the author never wrote -- and
# the only place that was ever said was a render manifest nobody reads before
# exporting. It is an error, for the same reason image_path is one: the output
# is wrong in a way the editor does not show.
# ---------------------------------------------------------------------------


def _ref_project(blocks, extra=None):
    outline = [{"id": "n1", "title": "Results", "blocks": blocks, "children": []}]
    if extra:
        outline.append(extra)
    return {"schema_version": 1, "outline": outline}


def test_dangling_reference_is_an_error():
    para = {"type": "para", "id": "p1",
            "runs": [{"t": "The spread is shown in "}, {"ref": "fig-1"}, {"t": "."}]}
    figure = {"type": "image", "id": "fig-1", "file": "images/spread.png",
              "caption": "Corner spread"}

    live = _lint_codes(_ref_project([para, figure]), _FULL_CFG)
    check("dangling_ref" not in live,
          "a reference to a figure that is there stays silent", "codes=%r" % live)

    gone = cl.lint_project(_ref_project([para]), _FULL_CFG)
    codes = [f["code"] for f in gone.flat]
    check("dangling_ref" in codes,
          "deleting the figure reports the sentence that pointed at it",
          "codes=%r" % codes)
    hits = [f for f in gone.flat if f["code"] == "dangling_ref"]
    check(len(hits) == 1, "once, for the one broken run", "%d hits" % len(hits))
    check(hits and hits[0]["level"] == "error", "as an error, not a warning",
          "%r" % (hits and hits[0].get("level")))
    check(gone.has_errors, "so it stands in the way of a package / fill workflow")
    check(hits and hits[0]["blockId"] == "p1" and hits[0]["nodeId"] == "n1",
          "and the Check tab can jump to the paragraph holding it", "%r" % hits)
    check(cl.LEVELS["dangling_ref"] == "error",
          "the renderer's own manifest warning is classified the same way")

    # A reference is resolved against the WHOLE document, so pointing forward or
    # into another section is not a mistake.
    later = _ref_project([para], extra={"id": "n2", "title": "Appendix",
                                        "blocks": [figure], "children": []})
    check("dangling_ref" not in _lint_codes(later, _FULL_CFG),
          "a reference into a later section still resolves")

    # What the engine bookmarks is what resolves (engine._collect_ref_targets):
    # an uncaptioned FIGURE is a target, an uncaptioned TABLE is not.
    check("dangling_ref" not in _lint_codes(
        _ref_project([para, dict(figure, caption="")]), _FULL_CFG),
        "an uncaptioned figure is still a target")
    table = {"type": "table", "id": "fig-1", "rows": [["a"]], "caption": ""}
    check("dangling_ref" in _lint_codes(_ref_project([para, table]), _FULL_CFG),
          "an uncaptioned table is not one")

    # A fixed template body contributes none of its own blocks.
    fixed = {"id": "n2", "title": "Boilerplate", "fixed_body": True,
             "blocks": [figure], "children": []}
    check("dangling_ref" in _lint_codes(_ref_project([para], extra=fixed), _FULL_CFG),
          "a figure inside a fixed template body is not a target either")


# ---------------------------------------------------------------------------
# 4b. A reference column is not this stage's work.
#
# An inherited report keeps the previous stage's numbers in a read-only
# reference column beside its own, still-empty cells. The pre-fix rule read
# EVERY entry in row["sims"], so those borrowed numbers made every unfilled row
# look done: zero warnings for exactly the table that carries all the work.
# ---------------------------------------------------------------------------

_REF_KEY = "ref_cdr_sim"


def _ref_group(read_only=True, role="reference"):
    g = {"key": _REF_KEY, "title": "Earlier stage", "stage": "CDR",
         "axes": ["MIN", "TYP", "MAX"],
         "source": {"report": "1108/CLKDIV_5G/CDR", "stage": "CDR",
                    "block": "b", "group": "sim", "hash": "fixture"}}
    if role is not None:
        g["role"] = role
    if read_only is not None:
        g["readOnly"] = read_only
    return g


def _own_group():
    return {"key": "own", "title": "This stage", "stage": "FDR",
            "axes": ["MIN", "TYP", "MAX"]}


def _inherited_row(item, own=None, ref=(1.0, 2.0, 3.0)):
    """A result row as report creation leaves it: this stage's cells cleared,
    the reference column carrying the earlier stage's numbers."""
    sims = {}
    if own is not None:
        sims["own"] = {"mtm": list(own), "ntwc": None}
    if ref is not None:
        sims[_REF_KEY] = {"mtm": list(ref), "ntwc": ref[2]}
    return {"cat": "Supply", "item": item, "kind": "result", "unit": "uA",
            "limit": "le", "spec_mtm": [None, None, 500],
            "sim_mtm": [None, None, None], "sim_ntwc": None, "sims": sims}


def _inherited_project(rows, sims=None):
    return _one_block_project(_dt(
        [dict(_GOOD_SETTING)] + rows,
        sims=sims if sims is not None else [_own_group(), _ref_group()]))


def _empty_sim_findings(project, cfg=None):
    return [f for f in cl.lint_project(project, cfg or _FULL_CFG).flat
            if f["code"] == "empty_sim_result"]


def test_reference_column_is_not_this_stage_work():
    # --- the defect, as an assertion on the reading itself ------------------
    row = _inherited_row("I_dd")
    old_rule = cl._row_sim_values(row)                  # every group counted
    check(not all(cl._is_empty(v) for v in old_rule),
          "the pre-fix reading of an inherited row looks filled in "
          "(this is the defect)", "%r" % old_rule)
    try:
        new_rule = cl._row_sim_values(row, {_REF_KEY})
    except TypeError as e:          # the pre-fix helper could not exclude a group
        new_rule = list(old_rule)
        check(False, "_row_sim_values accepts the groups to exclude", repr(e))
    check(all(cl._is_empty(v) for v in new_rule),
          "with the reference group excluded the same row reads as empty",
          "%r" % new_rule)

    # --- one warning per unfilled row --------------------------------------
    items = ["I_dd", "I_peak", "P_total"]
    proj = _inherited_project([_inherited_row(i) for i in items])
    found = _empty_sim_findings(proj)
    check(len(found) == len(items),
          "an inherited table warns once per unfilled result row "
          "(pre-fix: 0 of %d)" % len(items),
          "%d finding(s): %r" % (len(found), [f["message"] for f in found]))
    check(sorted(re.search(r'"(.*)"', f["message"]).group(1) for f in found)
          == sorted(items),
          "each warning names its own row",
          "%r" % [f["message"] for f in found])
    check(all(f["level"] == "warn" and f["blockId"] == "b" for f in found),
          "the warnings carry the block to jump to")

    # a setting row is judged by nobody, reference column or not
    check(all("Temperature" not in f["message"] for f in found),
          "a setting row is not reported as unfilled")

    # --- the opposite case: a row done for THIS stage stays silent ----------
    done = _inherited_project([_inherited_row("I_dd", own=(279, 412, 490))])
    check(_empty_sim_findings(done) == [],
          "a row with this stage's own values does not warn although a "
          "reference column sits beside it",
          "%r" % [f["message"] for f in _empty_sim_findings(done)])
    # ... including one that keeps its values in the flat schema, which is what
    # a table with no own group declared leaves behind
    flat_done = _inherited_project([dict(_inherited_row("I_dd"),
                                         sim_mtm=[279, 412, 490])])
    check(_empty_sim_findings(flat_done) == [],
          "flat sim_mtm counts as this stage's values",
          "%r" % [f["message"] for f in _empty_sim_findings(flat_done)])
    # a row whose own group is present but blank is still unfilled
    blank_own = _inherited_project([_inherited_row("I_dd",
                                                   own=(None, None, None))])
    check(len(_empty_sim_findings(blank_own)) == 1,
          "an own group full of blanks is still an unfilled row",
          "%r" % [f["message"] for f in _empty_sim_findings(blank_own)])

    # --- a reference group alone never satisfies the filled-in test ---------
    only_ref = _one_block_project(_dt(
        [dict(_GOOD_SETTING), _inherited_row("I_dd")], sims=[_ref_group()]))
    check(len(_empty_sim_findings(only_ref)) == 1,
          "a table whose only value group is a reference column is entirely "
          "unfilled",
          "%r" % [f["message"] for f in _empty_sim_findings(only_ref)])

    # both marks the creation path writes are honoured, separately
    for label, group in (("role only", _ref_group(read_only=None)),
                         ("readOnly only", _ref_group(role=None))):
        p = _inherited_project([_inherited_row("I_dd")],
                               sims=[_own_group(), group])
        check(len(_empty_sim_findings(p)) == 1,
              "a reference column marked by %s is excluded" % label,
              "%r" % [f["message"] for f in _empty_sim_findings(p)])

    # --- a plain second simulation column is NOT a reference ----------------
    two_sims = _inherited_project(
        [dict(_inherited_row("I_dd", ref=None),
              sims={"post": {"mtm": [279, 412, 490], "ntwc": None}})],
        sims=[_own_group(), {"key": "post", "title": "Post-layout",
                             "axes": ["MIN", "TYP", "MAX"]}])
    check(_empty_sim_findings(two_sims) == [],
          "a second ordinary simulation group still counts as filled in",
          "%r" % [f["message"] for f in _empty_sim_findings(two_sims)])

    # --- a genuinely empty table still warns, reference column or not -------
    no_ref = _one_block_project(_dt([dict(_GOOD_SETTING),
                                     _row(sim_mtm=[None, None, None],
                                          sim_ntwc=None)]))
    check(len(_empty_sim_findings(no_ref)) == 1,
          "an empty table with no reference column warns as before",
          "%r" % [f["message"] for f in _empty_sim_findings(no_ref)])

    # nothing else changed: the inherited table produces the unfilled-row
    # warnings and nothing new
    codes = sorted({f["code"] for f in cl.lint_project(proj, _FULL_CFG).flat})
    check(codes == ["empty_sim_result"],
          "the inherited fixture produces exactly the unfilled-row warnings",
          "%r" % codes)


# ---------------------------------------------------------------------------
# 5. CLI exit codes.
# ---------------------------------------------------------------------------


def _write_project(root, name, project):
    d = os.path.join(root, name)
    os.makedirs(os.path.join(d, "images"), exist_ok=True)
    with open(os.path.join(d, "project.json"), "w", encoding="utf-8") as fh:
        json.dump(project, fh)
    return d


def _run_cli(argv):
    """Return (exit_code, stdout)."""
    out = io.StringIO()
    err = io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = cl.main(argv)
    return rc, out.getvalue()


def _run_cli_full(argv):
    """Return (exit_code, stdout, stderr) -- for cases that check the error text."""
    out = io.StringIO()
    err = io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = cl.main(argv)
    return rc, out.getvalue(), err.getvalue()


def test_bom_project_loads():
    """A project.json (or template config) saved with a leading UTF-8 BOM -- a
    common artefact from Windows editors / Excel / Notepad -- must lint cleanly
    instead of crashing with an uncaught json.decoder.JSONDecodeError."""
    tmp = tempfile.mkdtemp(prefix="lint_bom_")
    try:
        clean_dir = os.path.join(tmp, "clean")
        os.makedirs(clean_dir, exist_ok=True)
        project = {"schema_version": 1, "outline": [
            {"id": "n", "title": "MODULE_A", "blocks": [
                {"type": "para", "runs": [{"t": "text"}]},
                _dt(_clean_table_rows())], "children": []}]}
        with open(os.path.join(clean_dir, "project.json"), "wb") as fh:
            fh.write(b"\xef\xbb\xbf" + json.dumps(project).encode("utf-8"))
        cfg_path = os.path.join(tmp, "cfg.json")
        with open(cfg_path, "wb") as fh:
            fh.write(b"\xef\xbb\xbf" + json.dumps(_FULL_CFG).encode("utf-8"))

        rc, out = _run_cli([clean_dir, "--config", cfg_path])
        check(rc == 0, "a BOM'd project.json + config lints clean (was a crash)",
              "rc=%r out=%r" % (rc, out))
        check("0 error" in out, "the BOM did not stop the findings from coming through",
              out)

        broken_dir = os.path.join(tmp, "broken")
        os.makedirs(broken_dir, exist_ok=True)
        with open(os.path.join(broken_dir, "project.json"), "wb") as fh:
            fh.write(b"\xef\xbb\xbf{not valid json")
        rc, out, err = _run_cli_full([broken_dir])
        check(rc == 2, "genuinely invalid JSON still exits 2, not a traceback",
              "rc=%r" % rc)
        check("project.json" in err and "Traceback" not in err,
              "the error names the file instead of dumping a traceback", err)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_cli_exit_codes():
    """0 = nothing to fix, 1 = at least one error (this is what gates a package),
    2 = there is no project.json to lint. Unchanged by the v2 rules: the new
    warnings and notes must NOT start failing a build."""
    tmp = tempfile.mkdtemp(prefix="lint_cli_")
    try:
        clean_dir = _write_project(tmp, "clean", {"schema_version": 1, "outline": [
            {"id": "n", "title": "MODULE_A", "blocks": [
                {"type": "para", "runs": [{"t": "text"}]},
                _dt(_clean_table_rows())], "children": []}]})
        # warnings + notes only: no setting rows, a blank limit, a cleared frame
        warn_dir = _write_project(tmp, "warned", {"schema_version": 1, "outline": [
            {"id": "n", "title": "TXBUF_2G", "blocks": [
                {"type": "image", "id": "i", "file": "", "caption": "Layout"},
                _dt([_row(limit=None)])], "children": []}]})
        err_dir = _write_project(tmp, "broken", {"schema_version": 1, "outline": [
            {"id": "n", "title": "TXBUF_2G", "blocks": [
                {"type": "datatable", "id": "b"}], "children": []}]})
        missing_dir = os.path.join(tmp, "nothing_here")
        os.makedirs(missing_dir)
        cfg_path = os.path.join(tmp, "cfg.json")
        with open(cfg_path, "w", encoding="utf-8") as fh:
            json.dump(_FULL_CFG, fh)

        rc, out = _run_cli([clean_dir, "--config", cfg_path])
        check(rc == 0, "clean project -> exit 0", "rc=%r" % rc)
        check("0 error" in out, "clean project prints a zero-error tally", out)

        rc, out = _run_cli([warn_dir, "--config", cfg_path])
        check(rc == 0, "warnings and notes alone -> exit 0 (they do not gate)",
              "rc=%r out=%r" % (rc, out))
        check("image_placeholder" not in out,
              "info findings are hidden without --info", out)
        rc, out = _run_cli([warn_dir, "--config", cfg_path, "--info"])
        check(rc == 0, "--info does not change the exit code", "rc=%r" % rc)
        check("image_placeholder" in out, "--info shows the notes", out)

        rc, out = _run_cli([err_dir, "--config", cfg_path])
        check(rc == 1, "an error -> exit 1", "rc=%r out=%r" % (rc, out))
        rc, out = _run_cli([err_dir, "--config", cfg_path, "--json"])
        check(rc == 1, "--json keeps exit 1", "rc=%r" % rc)
        parsed = json.loads(out)
        check(isinstance(parsed, list) and parsed[0]["code"] == "datatable_no_data",
              "--json prints the flat findings list", out[:200])
        rc, out = _run_cli([err_dir, "--config", cfg_path, "--grouped"])
        check(rc == 1, "--grouped keeps exit 1", "rc=%r" % rc)
        grouped = json.loads(out)
        check(sorted(grouped) == ["errors", "notes", "warnings"],
              "--grouped prints the Check-tab object", out[:200])
        rc, out = _run_cli([clean_dir, "--config", cfg_path, "--grouped"])
        check(rc == 0, "--grouped on a clean project -> exit 0", "rc=%r" % rc)

        rc, _ = _run_cli([missing_dir])
        check(rc == 2, "no project.json -> exit 2", "rc=%r" % rc)

        # --files adds the on-disk picture check; a missing picture is a warning,
        # so it must not start failing a build either.
        files_dir = _write_project(tmp, "files", {"schema_version": 1, "outline": [
            {"id": "n", "title": "MODULE_A", "blocks": [
                {"type": "image", "id": "i", "file": "images/gone.png",
                 "caption": "Layout"}], "children": []}]})
        rc, out = _run_cli([files_dir, "--config", cfg_path, "--files"])
        check(rc == 0, "--files with a missing picture -> still exit 0", "rc=%r" % rc)
        check("missing_image" in out, "--files reports the missing picture", out)
        rc, out = _run_cli([files_dir, "--config", cfg_path])
        check("missing_image" not in out,
              "without --files the on-disk check does not run", out)

        # the CLI runs without a config at all (a fill script's default path)
        rc, _ = _run_cli([clean_dir])
        check(rc == 0, "no --config -> still exit 0 on a clean project", "rc=%r" % rc)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main():
    report = cl.lint_project(FIXTURE, {})
    got = test_coverage(report)
    test_item_shape(report, got)
    test_every_item_carries_loc_node_message(report)
    test_grouping(report)
    test_sim_span_matches_renderer()
    test_dead_codes_are_gone()
    test_sim_span_mirrors_the_real_renderer()
    test_each_rule_fires_and_stays_silent()
    test_inherited_report_full_of_empty_frames_is_not_broken()
    test_dangling_reference_is_an_error()
    test_reference_column_is_not_this_stage_work()
    test_missing_image_needs_a_folder()
    test_cli_exit_codes()
    test_bom_project_loads()
    test_helpers_and_edge_cases()

    print("\n%d finding(s) from fixture; %d test failure(s)%s"
          % (len(report.flat), fails,
             "; %d skipped (%s)" % (len(skipped), ", ".join(skipped))
             if skipped else ""))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
