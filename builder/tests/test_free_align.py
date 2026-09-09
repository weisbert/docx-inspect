#!/usr/bin/env python3
"""A plain table's colour and its alignment survive the trip into Word.

Two faults, one table, both invisible to anyone reading only the editor:

  1. ROW KINDS NEVER REACHED THE DOCUMENT. A row's kind is what colours it --
     it is why a colour travels with a row instead of with a row NUMBER, and
     migrate_row_kind.py exists to write the field. engine.py rendered every
     free table without passing ``row_kinds``, so the renderer saw no kinds at
     all and fell back to the legacy index-keyed ``row_fills``. The editor drew
     the kind's colour, Word printed the legacy one, and nothing anywhere said
     the two disagreed. A table carrying kinds and no legacy map -- which is
     every table written since -- printed with no shading whatsoever.

  2. A CELL COULD NOT BE ALIGNED. Every free-table cell was centred, with the
     alignment written into the renderer as a constant. A column of running
     text (a numbered list of conditions, say) read as badly centred prose
     while the short labels beside it wanted the centre they had.

WHAT THIS FILE ASSERTS
  1. kinds alone shade the document, through a REAL engine.render_report
  2. kinds beat a legacy fill on the same row, in the document
  3. a table with only row_fills is untouched (the migration keeps both fields,
     so this is the shape most reports on disk still have)
  4. col_align aligns a column; a cell's own ``align`` overrides its column
  5. a table that asks for no alignment is centred exactly as before
  6. the .xlsx export says the same thing about the same table -- it is sold as
     a visual copy of the Word table, and it was reading neither field

AGAINST THE PRE-FIX FILES: 1, 2, 4 and 6 fail (no shading, no alignment).

The fixture is ASCII-only and belongs to this file. The template config comes
from test_render_golden.golden_config() so both tests describe the same
document; nothing here writes to a real report.

Run:
    python builder/tests/test_free_align.py
"""

import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
BUILDER = os.path.dirname(HERE)
if BUILDER not in sys.path:
    sys.path.insert(0, BUILDER)
import buildpath  # noqa: E402,F401  (registers core/docx_io/store/sync/web)

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH as ALIGN  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

import engine  # noqa: E402
import xlsx_export  # noqa: E402
from test_render_golden import golden_config  # noqa: E402

SETTING_FILL = "EEECE1"          # tables.py's default fill for a 'setting' row
HEADER_FILL = "D9D9D9"           # golden_config()'s free_table header_fill
LEGACY_FILL = "FFF2CC"           # a colour no kind maps to

FAILED = []


def check(ok, label, detail=""):
    print(("  PASS  " if ok else "  FAIL  ") + label + (("  [%s]" % detail) if detail else ""))
    if not ok:
        FAILED.append(label)


# ---------------------------------------------------------------------------
# The tables under test.
# ---------------------------------------------------------------------------
def kinds_only_table():
    """What the editor writes today: kinds, and no legacy colour map at all."""
    return {
        "type": "table", "id": "tbl-kinds", "caption": "Conditions and results",
        "header_rows": 1,
        "rows": [
            ["Item", "Value", "Unit"],
            ["Supply", "1.1", "V"],
            ["Temperature", "25", "C"],
            ["Gain margin", "11.4", "dB"],
        ],
        "row_kinds": ["header", "setting", "setting", "result"],
    }


def kind_over_legacy_table():
    """Both fields, disagreeing: the kind decides, in the editor and in Word."""
    return {
        "type": "table", "id": "tbl-both", "caption": "Both fields",
        "header_rows": 1,
        "rows": [
            ["Part", "Reading"],
            ["Supply", "1.1"],
            ["Gain margin", "11.4"],
        ],
        "row_kinds": ["header", "setting", "result"],
        # row 2 is a result: the renderer must leave it unshaded, NOT paint it
        # with the stale colour this map still carries for it.
        "row_fills": {"1": LEGACY_FILL, "2": LEGACY_FILL},
    }


def legacy_only_table():
    """The shape most reports on disk still have: colours addressed by index."""
    return {
        "type": "table", "id": "tbl-legacy", "caption": "Legacy colours",
        "header_rows": 1,
        "rows": [
            ["Device", "Reading"],
            ["Supply", "1.1"],
            ["Gain margin", "11.4"],
        ],
        "row_fills": {"1": LEGACY_FILL},
    }


def aligned_table():
    """A column of running text reads left; the labels beside it stay centred;
    one cell overrules its column."""
    return {
        "type": "table", "id": "tbl-align", "caption": "Alignment",
        "header_rows": 1,
        "rows": [
            ["Check", "Finding"],
            ["EOS", {"runs": [{"t": "1. no violation\n2. within the limit"}]}],
            ["Margin", {"runs": [{"t": "see above"}], "align": "right"}],
            ["Note", "plain text, no runs"],
        ],
        "col_align": ["center", "left"],
    }


def plain_table():
    """No kinds, no alignment: the control. Must render as it always did."""
    return {
        "type": "table", "id": "tbl-plain", "caption": "Plain",
        "header_rows": 1,
        "rows": [["Name", "Number"], ["Supply", "1.1"]],
    }


def project():
    return {
        "schema_version": 1,
        "template": "golden_tpl_v1",
        "meta": {"title": "Free table fields", "author": "Tester",
                 "reviewers": [], "revisions": []},
        "outline": [
            {"id": "ch1", "title": "Tables", "blocks": [
                kinds_only_table(), kind_over_legacy_table(), legacy_only_table(),
                aligned_table(), plain_table(),
            ], "children": []},
        ],
    }


# ---------------------------------------------------------------------------
# Reading the rendered document.
# ---------------------------------------------------------------------------
def cell_fill(cell):
    """The cell's shading colour, or None when it carries no shading."""
    tcPr = cell._tc.find(qn("w:tcPr"))
    if tcPr is None:
        return None
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        return None
    val = shd.get(qn("w:fill"))
    return None if val in (None, "auto") else val.upper()


def row_fill(table, r):
    return cell_fill(table.rows[r].cells[0])


def cell_align(cell):
    return cell.paragraphs[0].alignment


def text_of(cell):
    return "\n".join(p.text for p in cell.paragraphs).strip()


def table_by_first_cell(doc, header):
    """The rendered table whose header row reads exactly ``header``."""
    for t in doc.tables:
        if t.rows and [text_of(c) for c in t.rows[0].cells] == header:
            return t
    return None


def render(tmp, font_pt=None):
    """Render the fixture. ``font_pt`` is the free_table config's own size: the
    real templates set one (8pt), and only a sized cell has ever had its
    alignment written at all -- an unsized one is left to the style. Both are
    rendered here so the alignment work can be shown not to have moved either."""
    cfg = golden_config()
    cfg["_logo_path"] = ""
    name = "sized"
    if font_pt:
        cfg["free_table"] = dict(cfg["free_table"], font_pt=font_pt)
    else:
        name = "unsized"
    out = os.path.join(tmp, "out", "free_align_%s.docx" % name)
    engine.render_report(project(), cfg, tmp, out)
    return Document(out)


# ---------------------------------------------------------------------------
# 1-3. kinds reach the document
# ---------------------------------------------------------------------------
def test_kinds_reach_word(doc):
    print("== a row's kind colours it in the document ==")
    t = table_by_first_cell(doc, ["Item", "Value", "Unit"])
    check(t is not None, "the kinds-only table was rendered")
    if t is None:
        return
    check(row_fill(t, 0) == HEADER_FILL, "the header row takes the header fill",
          str(row_fill(t, 0)))
    check(row_fill(t, 1) == SETTING_FILL,
          "a 'setting' row is shaded from its KIND, with no row_fills anywhere",
          str(row_fill(t, 1)))
    check(row_fill(t, 2) == SETTING_FILL, "the second condition row too",
          str(row_fill(t, 2)))
    check(row_fill(t, 3) is None, "a 'result' row is left unshaded",
          str(row_fill(t, 3)))
    # the shading is on the WHOLE row, not just its first cell
    check(all(cell_fill(c) == SETTING_FILL for c in t.rows[1].cells),
          "every cell of the condition row carries the fill")


def test_kind_beats_legacy_fill(doc):
    print("== a kind overrules the legacy colour on the same row ==")
    t = table_by_first_cell(doc, ["Part", "Reading"])
    check(t is not None, "the kind+row_fills table was rendered")
    if t is None:
        return
    check(row_fill(t, 1) == SETTING_FILL,
          "the condition row takes the kind's colour, not the stale map's",
          str(row_fill(t, 1)))
    check(row_fill(t, 2) is None,
          "the result row is unshaded even though the map still names a colour",
          str(row_fill(t, 2)))


def test_legacy_only_unchanged(doc):
    print("== a table with only row_fills is untouched ==")
    t = table_by_first_cell(doc, ["Device", "Reading"])
    check(t is not None, "the legacy table was rendered")
    if t is None:
        return
    check(row_fill(t, 1) == LEGACY_FILL, "it kept its index-keyed colour",
          str(row_fill(t, 1)))
    check(row_fill(t, 0) == HEADER_FILL, "its header row is unchanged")
    check(row_fill(t, 2) is None, "the row the map does not name is unshaded")
    check(all(cell_align(c) == ALIGN.CENTER for c in t.rows[1].cells),
          "and every cell of it is still centred")


# ---------------------------------------------------------------------------
# 4-5. alignment
# ---------------------------------------------------------------------------
def test_alignment(doc):
    print("== a column, and a cell, can read left or right ==")
    t = table_by_first_cell(doc, ["Check", "Finding"])
    check(t is not None, "the aligned table was rendered")
    if t is None:
        return
    check(cell_align(t.rows[0].cells[0]) == ALIGN.CENTER,
          "the first column is centred, as its col_align says")
    check(cell_align(t.rows[1].cells[1]) == ALIGN.LEFT,
          "the running-text column reads left",
          str(cell_align(t.rows[1].cells[1])))
    check(cell_align(t.rows[2].cells[1]) == ALIGN.RIGHT,
          "a cell's own align overrules its column",
          str(cell_align(t.rows[2].cells[1])))
    check(cell_align(t.rows[3].cells[1]) == ALIGN.LEFT,
          "a plain-text cell follows its column too",
          str(cell_align(t.rows[3].cells[1])))
    check("2. within the limit" in text_of(t.rows[1].cells[1]),
          "a line break inside a cell survives")
    check(cell_align(t.rows[1].cells[0]) == ALIGN.CENTER,
          "the labels beside it keep the centre")


def test_no_alignment_is_centred(doc):
    print("== a table that asks for nothing is centred, as before ==")
    t = table_by_first_cell(doc, ["Name", "Number"])
    check(t is not None, "the plain table was rendered")
    if t is None:
        return
    check(all(cell_align(c) == ALIGN.CENTER for row in t.rows for c in row.cells),
          "every cell is centred")
    check(all(cell_fill(c) is None for c in t.rows[1].cells),
          "and its body row carries no shading")


# ---------------------------------------------------------------------------
# 6. the spreadsheet export says the same thing
# ---------------------------------------------------------------------------
def test_xlsx_mirrors(tmp):
    print("== the .xlsx export reads the same two fields ==")
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("  SKIP  openpyxl not installed")
        return
    cfg = golden_config()

    path = os.path.join(tmp, "kinds.xlsx")
    with open(path, "wb") as fh:
        fh.write(xlsx_export.build_block_xlsx(kinds_only_table(), cfg))
    ws = load_workbook(path).active
    check(ws.cell(row=2, column=1).fill.fgColor.rgb == "FF" + SETTING_FILL,
          "a condition row is shaded in the sheet from its kind",
          str(ws.cell(row=2, column=1).fill.fgColor.rgb))
    check(ws.cell(row=4, column=1).fill.fill_type in (None, "none"),
          "a result row is left unshaded",
          str(ws.cell(row=4, column=1).fill.fill_type))

    path = os.path.join(tmp, "align.xlsx")
    with open(path, "wb") as fh:
        fh.write(xlsx_export.build_block_xlsx(aligned_table(), cfg))
    ws = load_workbook(path).active
    check(ws.cell(row=2, column=2).alignment.horizontal == "left",
          "the running-text column reads left in the sheet",
          str(ws.cell(row=2, column=2).alignment.horizontal))
    check(ws.cell(row=2, column=2).alignment.wrap_text is True,
          "a cell with its own line breaks wraps (Excel does not on its own)")
    check(ws.cell(row=3, column=2).alignment.horizontal == "right",
          "a cell's own align overrules its column here too")
    check(ws.cell(row=2, column=1).alignment.horizontal == "center",
          "the centred column is still centred")

    path = os.path.join(tmp, "plain.xlsx")
    with open(path, "wb") as fh:
        fh.write(xlsx_export.build_block_xlsx(plain_table(), cfg))
    ws = load_workbook(path).active
    check(ws.cell(row=2, column=1).alignment.horizontal == "center",
          "a table that asks for nothing is centred in the sheet too")


def test_unsized_config_is_untouched(doc):
    """A free_table config with no font_pt never wrote an alignment, and still
    does not -- until a table names one."""
    print("== an unsized table is left where it was ==")
    t = table_by_first_cell(doc, ["Name", "Number"])
    check(t is not None, "the plain table was rendered")
    if t is not None:
        check(all(cell_align(c) is None for row in t.rows for c in row.cells),
              "no alignment is written at all, as before")
    a = table_by_first_cell(doc, ["Check", "Finding"])
    check(a is not None, "the aligned table was rendered")
    if a is not None:
        check(cell_align(a.rows[1].cells[1]) == ALIGN.LEFT,
              "but a named alignment is honoured whether or not the config sizes the text",
              str(cell_align(a.rows[1].cells[1])))
        check(cell_align(a.rows[2].cells[1]) == ALIGN.RIGHT,
              "including a cell's own")


def main():
    tmp = tempfile.mkdtemp(prefix="free_align_")
    os.makedirs(os.path.join(tmp, "images"), exist_ok=True)
    doc = render(tmp, font_pt=8)
    test_kinds_reach_word(doc)
    test_kind_beats_legacy_fill(doc)
    test_legacy_only_unchanged(doc)
    test_alignment(doc)
    test_no_alignment_is_centred(doc)
    test_unsized_config_is_untouched(render(tmp))
    test_xlsx_mirrors(tmp)
    print("")
    if FAILED:
        print("FAILED (%d): %s" % (len(FAILED), ", ".join(FAILED)))
        return 1
    print("all free-table colour and alignment assertions passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
